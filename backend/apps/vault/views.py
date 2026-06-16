import openpyxl
from io import BytesIO
from django.http import HttpResponse
from django.db import transaction
from django.db.models import Q
from django.utils import timezone
from rest_framework import viewsets, permissions, status
from rest_framework.decorators import action
from rest_framework.response import Response

from apps.vault.models import PasswordVault
from apps.organizations.models import Organization
from apps.vault.serializers import PasswordVaultSerializer
from apps.authentication.rbac import HasGranularPermission
from apps.audit.models import AuditLog

class PasswordVaultViewSet(viewsets.ModelViewSet):
    """
    ViewSet managing shared server/FTP passwords.
    Integrates live search, granular RBAC, dynamic audits, and Excel integrations.
    """
    queryset = PasswordVault.objects.all().select_related('organization', 'created_by', 'modified_by').order_by('-created_at')
    serializer_class = PasswordVaultSerializer
    permission_classes = [permissions.IsAuthenticated, HasGranularPermission]
    required_permission = 'vault:view'

    def get_queryset(self):
        queryset = PasswordVault.objects.all().select_related('organization', 'created_by', 'modified_by').order_by('-created_at')
        
        # Search by Title, Login, associated email, or Organization
        search = self.request.query_params.get('search', None)
        if search:
            queryset = queryset.filter(
                Q(title__icontains=search) |
                Q(login_username__icontains=search) |
                Q(associated_email__icontains=search) |
                Q(organization__name__icontains=search)
            )

        # Filter by organization
        org_filter = self.request.query_params.get('organization', None)
        if org_filter:
            queryset = queryset.filter(organization_id=org_filter)

        return queryset

    def perform_create(self, serializer):
        vault_item = serializer.save()
        # Log creation
        self._log_audit_event(
            action='create_password',
            details={
                'password_id': str(vault_item.id),
                'title': vault_item.title,
                'organization': vault_item.organization.name,
                'login_username': vault_item.login_username
            }
        )

    def perform_update(self, serializer):
        vault_item = serializer.save()
        # Log update
        self._log_audit_event(
            action='update_password',
            details={
                'password_id': str(vault_item.id),
                'title': vault_item.title,
                'organization': vault_item.organization.name,
                'login_username': vault_item.login_username
            }
        )

    def perform_destroy(self, instance):
        # Log delete
        self._log_audit_event(
            action='delete_password',
            details={
                'password_id': str(instance.id),
                'title': instance.title,
                'organization': instance.organization.name,
                'login_username': instance.login_username
            }
        )
        instance.delete()

    @action(detail=True, methods=['get'], url_path='reveal')
    def reveal_password(self, request, pk=None):
        """
        Endpoint to retrieve the plain text password (reveals the password on screen).
        Enforces 'vault:reveal' permission and creates an audit entry.
        """
        if not request.user.has_permission('vault:reveal'):
            return Response({"detail": "Nu aveți permisiunea de a dezvălui parole."}, status=status.HTTP_403_FORBIDDEN)
            
        instance = self.get_object()
        plain_password = instance.get_password()
        
        # Log reveal event
        self._log_audit_event(
            action='reveal_password',
            details={
                'password_id': str(instance.id),
                'title': instance.title,
                'organization': instance.organization.name,
                'login_username': instance.login_username
            }
        )
        
        return Response({"password": plain_password})

    @action(detail=True, methods=['get'], url_path='copy')
    def copy_password(self, request, pk=None):
        """
        Endpoint to retrieve plain text password for copy-to-clipboard action.
        Enforces 'vault:copy' permission and creates an audit entry.
        """
        if not request.user.has_permission('vault:copy'):
            return Response({"detail": "Nu aveți permisiunea de a copia parole."}, status=status.HTTP_403_FORBIDDEN)
            
        instance = self.get_object()
        plain_password = instance.get_password()
        
        # Log copy event
        self._log_audit_event(
            action='copy_password',
            details={
                'password_id': str(instance.id),
                'title': instance.title,
                'organization': instance.organization.name,
                'login_username': instance.login_username
            }
        )
        
        return Response({"password": plain_password})

    def _normalize_header(self, h):
        if not h:
            return ""
        # Split by any whitespace (including non-breaking spaces \xa0) and join with regular space
        h = " ".join(str(h).split()).lower()
        replacements = {
            'ă': 'a', 'â': 'a', 'î': 'i', 'ș': 's', 'ț': 't',
            'ş': 's', 'ţ': 't', '/': ' ', '-': ' ', '_': ' '
        }
        for k, v in replacements.items():
            h = h.replace(k, v)
        return " ".join(h.split())

    @action(detail=False, methods=['post'], url_path='import')
    def import_excel(self, request):
        """
        Import passwords from Excel.
        Supports flexible column headers and diacritics.
        """
        if not request.user.has_permission('vault:import'):
            return Response({"detail": "Nu aveți permisiunea de a importa parole."}, status=status.HTTP_403_FORBIDDEN)

        excel_file = request.FILES.get('file')
        if not excel_file:
            return Response({"detail": "Niciun fișier nu a fost trimis."}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
            ws = wb.active
            
            imported_count = 0
            skipped_count = 0
            errors = []
            
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return Response({"detail": "Fișierul este gol."}, status=status.HTTP_400_BAD_REQUEST)
                
            headers = [self._normalize_header(h) for h in rows[0]]
            
            print(f"IMPORT DEBUG: Active Sheet Title: {ws.title}")
            print(f"IMPORT DEBUG: Total Rows Detected: {len(rows)}")
            print(f"IMPORT DEBUG: Normalized Headers: {headers}")
            
            # Flexible header mapping
            org_keys = ['organizatia', 'organizatie', 'org']
            title_keys = ['redenumita', 'denumirea noua', 'denumire noua', 'denumire nou', 'denumirea nou', 'denumire ip', 'denumire', 'title', 'nume']
            login_keys = ['login', 'login user', 'username', 'utilizator', 'user']
            pass_keys = ['parola', 'parole', 'password', 'pass']
            email_keys = ['email', 'asociat email', 'mail']
            phone_keys = ['telefon', 'telefon asociat', 'nr telefon', 'numar telefon', 'phone']

            def find_header_index(keys):
                for key in keys:
                    norm = self._normalize_header(key)
                    if norm in headers:
                        return headers.index(norm)
                return -1

            org_idx = find_header_index(org_keys)
            title_idx = find_header_index(title_keys)
            login_idx = find_header_index(login_keys)
            pass_idx = find_header_index(pass_keys)
            email_idx = find_header_index(email_keys)
            phone_idx = find_header_index(phone_keys)

            print(f"IMPORT DEBUG: Mapped Indexes -> Org: {org_idx}, Title: {title_idx}, Login: {login_idx}, Password: {pass_idx}, Email: {email_idx}, Phone: {phone_idx}")

            if org_idx == -1 or title_idx == -1 or login_idx == -1 or pass_idx == -1:
                missing = []
                if org_idx == -1: missing.append("Organizația")
                if title_idx == -1: missing.append("Redenumită (Denumire/IP)")
                if login_idx == -1: missing.append("Login")
                if pass_idx == -1: missing.append("Parola")
                return Response({"detail": f"Lipsește coloana necesară: {', '.join(missing)}"}, status=status.HTTP_400_BAD_REQUEST)

            seen_keys = set()
            with transaction.atomic():
                for row_num, row in enumerate(rows[1:], start=2):
                    if not any(row):  # Skip completely empty rows
                        continue
                    
                    org_name = str(row[org_idx]).strip() if org_idx < len(row) and row[org_idx] is not None else ""
                    title = str(row[title_idx]).strip() if title_idx < len(row) and row[title_idx] is not None else ""
                    login = str(row[login_idx]).strip() if login_idx < len(row) and row[login_idx] is not None else ""
                    password = str(row[pass_idx]).strip() if pass_idx < len(row) and row[pass_idx] is not None else ""
                    email = str(row[email_idx]).strip() if email_idx != -1 and email_idx < len(row) and row[email_idx] is not None else ""
                    phone = str(row[phone_idx]).strip() if phone_idx != -1 and phone_idx < len(row) and row[phone_idx] is not None else ""

                    if not org_name and not title and not login and not password:
                        continue

                    if not org_name or not login or not password:
                        reason = []
                        if not org_name: reason.append("Organizație")
                        if not login: reason.append("Login")
                        if not password: reason.append("Parolă")
                        err_msg = f"Rândul {row_num}: Date incomplete (Câmpurile {', '.join(reason)} sunt obligatorii)."
                        errors.append(err_msg)
                        skipped_count += 1
                        print(f"IMPORT DEBUG: Rândul {row_num} ignorat (date incomplete): {reason}. Row content: {row[:5]}")
                        continue

                    # Duplicate check within the uploaded Excel file
                    row_key = (org_name.lower(), title.lower(), login.lower())
                    if row_key in seen_keys:
                        err_msg = f"Rândul {row_num}: Credențial duplicat detectat în cadrul fișierului Excel ({org_name} - {title} - {login})."
                        errors.append(err_msg)
                        skipped_count += 1
                        print(f"IMPORT DEBUG: Rândul {row_num} ignorat (duplicat în fișier): {org_name} - {title} - {login}")
                        continue

                    # Clean organization name (normalize whitespace only, keeping underscores and dashes intact)
                    org_name_clean = " ".join(org_name.split())
                    org_code = org_name_clean.replace(' ', '_').upper()[:50]

                    # Find existing organization by name (case-insensitive) or code
                    org = Organization.objects.filter(
                        Q(name__iexact=org_name_clean) | Q(code=org_code)
                    ).first()

                    if not org:
                        try:
                            org = Organization.objects.create(
                                name=org_name_clean,
                                code=org_code
                            )
                        except Exception as org_err:
                            # Fallback if there is a racing issue or other unique constraint violation
                            org = Organization.objects.filter(code=org_code).first()
                            if not org:
                                raise org_err

                    duplicate_check = PasswordVault.objects.filter(
                        organization=org,
                        title=title,
                        login_username=login
                    ).exists()

                    if duplicate_check:
                        err_msg = f"Rândul {row_num}: Credențial existent în baza de date pentru această organizație ({org_name} - {title} - {login})."
                        errors.append(err_msg)
                        skipped_count += 1
                        print(f"IMPORT DEBUG: Rândul {row_num} ignorat (duplicat în DB): {org_name} - {title} - {login}")
                        seen_keys.add(row_key)
                        continue

                    vault_item = PasswordVault(
                        organization=org,
                        title=title,
                        login_username=login,
                        associated_email=email,
                        associated_phone=phone,
                        created_by=request.user
                    )
                    vault_item.set_password(password)
                    vault_item.save()
                    
                    seen_keys.add(row_key)
                    imported_count += 1

            self._log_audit_event(
                action='import_excel',
                details={
                    'imported_count': imported_count,
                    'skipped_count': skipped_count,
                    'file_name': excel_file.name
                }
            )

            return Response({
                "detail": f"Import finalizat. Adăugate: {imported_count}. Omise (duplicate/invalide): {skipped_count}.",
                "imported": imported_count,
                "skipped": skipped_count,
                "errors": errors
            }, status=status.HTTP_200_OK)

        except Exception as e:
            return Response({"detail": f"Eroare la procesarea fișierului Excel: {str(e)}"}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=False, methods=['get'], url_path='template')
    def download_template(self, request):
        """
        Generate and return a clean Excel template for password import.
        """
        if not request.user.has_permission('vault:import'):
            return Response({"detail": "Nu aveți permisiunea de a importa parole."}, status=status.HTTP_403_FORBIDDEN)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Import Template"
        
        # Headers
        headers = ["Organizația", "Redenumita", "Login", "Parola", "Email", "Telefon"]
        ws.append(headers)
        
        # Add a sample row
        ws.append(["SIDESI", "APC A0120-0238", "218_ccl", "ParolaSecreta123", "admin@ccl.ro", "0722123456"])
        
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = 'attachment; filename="SIDESI_Model_Import.xlsx"'
        return response

    @action(detail=False, methods=['get'], url_path='export')
    def export_excel(self, request):
        """
        Export all (or filtered) vault passwords in decrypted Excel sheet.
        """
        if not request.user.has_permission('vault:export'):
            return Response({"detail": "Nu aveți permisiunea de a exporta parole."}, status=status.HTTP_403_FORBIDDEN)

        # Apply same filters as queryset
        vault_items = self.get_queryset()

        # Build Excel Workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "SIDESI Vault Passwords"

        # Headers
        headers = ["Organizația", "Denumire Nou", "Login", "Parola", "Email", "Telefon", "Creat de", "Data creării"]
        ws.append(headers)

        # Fill Data
        for item in vault_items:
            try:
                decrypted_pass = item.get_password()
            except Exception:
                decrypted_pass = "Decryption_Error"
                
            ws.append([
                item.organization.name,
                item.title,
                item.login_username,
                decrypted_pass,
                item.associated_email or "",
                item.associated_phone or "",
                item.created_by.username,
                item.created_at.strftime('%Y-%m-%d %H:%M:%S')
            ])

        # Count before saving buffer
        exported_count = vault_items.count()

        # Save workbook to buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        # Log export event
        self._log_audit_event(
            action='export_excel',
            details={
                'exported_count': exported_count
            }
        )

        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = f'attachment; filename="SIDESI_Passwords_Export_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        return response

    def _log_audit_event(self, action, details):
        """Helper to create audit log records."""
        ip = getattr(self.request, 'client_ip', '0.0.0.0')
        ua = getattr(self.request, 'user_agent', 'system')
        
        AuditLog.objects.create(
            user=self.request.user,
            username_display=self.request.user.username,
            action=action,
            module='vault',
            ip_address=ip,
            user_agent=ua,
            details=details
        )
