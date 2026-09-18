import os
from io import BytesIO
from datetime import datetime

from django.conf import settings
from django.http import HttpResponse
from django.db.models import Q
from rest_framework import viewsets, permissions, status, parsers
from rest_framework.decorators import action
from rest_framework.response import Response

# Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.utils import get_column_letter

# PDF
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import cm, mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph,
    Spacer, Image as RLImage, HRFlowable
)
from reportlab.platypus.doctemplate import PageTemplate, BaseDocTemplate
from reportlab.platypus.frames import Frame
from reportlab.lib.colors import HexColor

from apps.authentication.rbac import HasGranularPermission
from apps.inventory.models import InventoryItem
from apps.inventory.serializers import InventoryItemSerializer
from apps.audit.models import AuditLog

# ── Brand colours ──────────────────────────────────────────────────────────────
SIDESI_DARK   = HexColor('#0F172A')  # Antet pagina
SIDESI_ACCENT = HexColor('#3B82F6')  # Titlu / accent
SIDESI_LIGHT  = HexColor('#EFF6FF')  # Fundal rânduri impare
SIDESI_MID    = HexColor('#DBEAFE')  # Fundal antet tabel
SIDESI_WHITE  = HexColor('#FFFFFF')
SIDESI_BORDER = HexColor('#CBD5E1')
SIDESI_TEXT   = HexColor('#1E293B')
SIDESI_MUTED  = HexColor('#64748B')


def _get_assigned_name(item):
    name = f"{item.assigned_to.first_name} {item.assigned_to.last_name}".strip()
    return name or item.assigned_to.username


def _format_dt(dt):
    return dt.strftime('%Y-%m-%d') if dt else '—'


# ─────────────────────────────────────────────────────────────────────────────
class InventoryItemViewSet(viewsets.ModelViewSet):
    """
    ViewSet for managing inventory objects.
    Supports CRUD, search by name/inventory_number, and filtering by assigned_to user.
    Image uploads are handled via multipart/form-data; Pillow compression is applied in model.save().
    """
    serializer_class = InventoryItemSerializer
    permission_classes = [permissions.IsAuthenticated, HasGranularPermission]
    required_permission = 'inventory:view'
    parser_classes = [parsers.MultiPartParser, parsers.FormParser, parsers.JSONParser]

    def get_queryset(self):
        qs = InventoryItem.objects.select_related('assigned_to').order_by('assigned_to__first_name', 'name')

        search = self.request.query_params.get('search', None)
        if search:
            qs = qs.filter(
                Q(name__icontains=search) |
                Q(inventory_number__icontains=search) |
                Q(description__icontains=search)
            )
        user_filter = self.request.query_params.get('assigned_to', None)
        if user_filter:
            qs = qs.filter(assigned_to__id=user_filter)

        return qs

    def get_permissions(self):
        if self.action in ['create', 'update', 'partial_update', 'destroy']:
            self.required_permission = 'inventory:manage'
        else:
            self.required_permission = 'inventory:view'
        return super().get_permissions()

    def perform_destroy(self, instance):
        self._log_audit('inventory_delete', {
            'item_id': str(instance.id),
            'name': instance.name,
            'inventory_number': instance.inventory_number,
        })
        instance.delete()

    def _log_audit(self, action, details):
        ip = getattr(self.request, 'client_ip', '0.0.0.0')
        ua = getattr(self.request, 'user_agent', 'system')[:512]
        AuditLog.objects.create(
            user=self.request.user,
            username_display=self.request.user.username,
            action=action,
            module='inventory',
            ip_address=ip,
            user_agent=ua,
            details=details
        )

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data, context={'request': request})
        serializer.is_valid(raise_exception=True)
        instance = serializer.save()
        output = InventoryItemSerializer(instance, context={'request': request})
        self._log_audit('inventory_create', {
            'item_id': str(instance.id),
            'name': instance.name,
            'inventory_number': instance.inventory_number,
        })
        return Response(output.data, status=status.HTTP_201_CREATED)

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        if 'image' not in request.FILES and 'image' not in request.data:
            data = request.data.copy()
        else:
            data = request.data
        serializer = self.get_serializer(instance, data=data, partial=partial, context={'request': request})
        serializer.is_valid(raise_exception=True)
        instance = serializer.save()
        output = InventoryItemSerializer(instance, context={'request': request})
        self._log_audit('inventory_update', {
            'item_id': str(instance.id),
            'name': instance.name,
            'fields_changed': list(request.data.keys()),
        })
        return Response(output.data)

    # ─────────────────────────────────────────────────────────────────────────
    #   EXCEL EXPORT
    # ─────────────────────────────────────────────────────────────────────────
    @action(detail=False, methods=['get'], url_path='export')
    def export_excel(self, request):
        """
        Export inventory as a styled Excel workbook.
        Query params:
          - assigned_to (UUID): filter by user
          - include_images (bool): embed photos in column A
        """
        queryset = self.filter_queryset(self.get_queryset())
        include_images = request.query_params.get('include_images', 'false').lower() == 'true'

        wb = Workbook()
        ws = wb.active
        ws.title = "Inventar SIDESI"

        # ── Print settings ──────────────────────────────────────────────────
        ws.page_setup.orientation  = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.paperSize    = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth   = 1
        ws.page_setup.fitToHeight  = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = '1:3'  # Repeat header on each printed page

        # ── Styles ──────────────────────────────────────────────────────────
        fn = 'Calibri'
        DARK   = '0F172A'
        ACCENT = '3B82F6'
        LIGHT  = 'EFF6FF'
        MID    = 'DBEAFE'
        ALT    = 'F8FAFC'

        title_font    = Font(name=fn, size=16, bold=True, color='FFFFFF')
        subtitle_font = Font(name=fn, size=9,  color='BFDBFE')
        header_font   = Font(name=fn, size=9,  bold=True, color='1E3A5F')
        data_font     = Font(name=fn, size=9,  color='1E293B')
        muted_font    = Font(name=fn, size=8,  italic=True, color='64748B')

        title_fill    = PatternFill(start_color=DARK,   end_color=DARK,   fill_type='solid')
        header_fill   = PatternFill(start_color=MID,    end_color=MID,    fill_type='solid')
        alt_fill      = PatternFill(start_color=LIGHT,  end_color=LIGHT,  fill_type='solid')
        white_fill    = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

        thin  = Side(style='thin',   color='CBD5E1')
        thick = Side(style='medium', color=ACCENT)
        cell_border   = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_border = Border(left=thick, right=thick, top=thick, bottom=thick)

        center = Alignment(horizontal='center', vertical='center', wrap_text=False)
        left   = Alignment(horizontal='left',   vertical='center', wrap_text=True)
        right  = Alignment(horizontal='right',  vertical='center')

        # ── Column definitions ───────────────────────────────────────────────
        if include_images:
            cols = [
                ('Poză',            'A', 14),
                ('Nr. Inventar',    'B', 18),
                ('Denumire Obiect', 'C', 28),
                ('Cantitate',       'D', 10),
                ('Responsabil',     'E', 22),
                ('Descriere',       'F', 36),
                ('Dată',           'G', 14),
            ]
        else:
            cols = [
                ('Nr. Inventar',    'A', 18),
                ('Denumire Obiect', 'B', 30),
                ('Cantitate',       'C', 10),
                ('Responsabil',     'D', 22),
                ('Descriere',       'E', 38),
                ('Dată Alocare',    'F', 14),
            ]

        num_cols = len(cols)
        last_col = cols[-1][1]

        # Apply column widths
        for _, col_letter, width in cols:
            ws.column_dimensions[col_letter].width = width

        # ── Row 1 — Title banner ─────────────────────────────────────────────
        ws.row_dimensions[1].height = 38
        ws.merge_cells(f'A1:{last_col}1')
        title_cell = ws['A1']
        title_cell.value    = '  CORPORAȚIA SIDESI  —  RAPORT INVENTAR'
        title_cell.font     = title_font
        title_cell.fill     = title_fill
        title_cell.alignment = center

        # ── Row 2 — Subtitle / metadata ──────────────────────────────────────
        ws.row_dimensions[2].height = 20
        ws.merge_cells(f'A2:{last_col}2')
        now = datetime.now().strftime('%d %B %Y, %H:%M')
        subtitle_cell = ws['A2']
        subtitle_cell.value     = f'  Generat la: {now}  |  Total: {queryset.count()} obiecte'
        subtitle_cell.font      = subtitle_font
        subtitle_cell.fill      = title_fill
        subtitle_cell.alignment = center

        # ── Row 3 — Column headers ───────────────────────────────────────────
        ws.row_dimensions[3].height = 22
        for col_idx, (label, col_letter, _) in enumerate(cols, start=1):
            cell = ws.cell(row=3, column=col_idx)
            cell.value     = label.upper()
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = center if col_idx in [1, 4] else left
            cell.border    = header_border

        # Enable AutoFilter on header row
        ws.auto_filter.ref = f'A3:{last_col}3'

        # ── Data rows ────────────────────────────────────────────────────────
        for row_idx, item in enumerate(queryset, start=4):
            is_odd  = (row_idx % 2 == 0)
            row_fill = alt_fill if is_odd else white_fill
            row_height = 55 if include_images else 22
            ws.row_dimensions[row_idx].height = row_height

            if include_images:
                row_vals = [
                    '',
                    item.inventory_number or '—',
                    item.name,
                    item.quantity,
                    _get_assigned_name(item),
                    item.description or '—',
                    _format_dt(item.created_at),
                ]
            else:
                row_vals = [
                    item.inventory_number or '—',
                    item.name,
                    item.quantity,
                    _get_assigned_name(item),
                    item.description or '—',
                    _format_dt(item.created_at),
                ]

            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value     = val
                cell.font      = data_font
                cell.fill      = row_fill
                cell.border    = cell_border
                cell.alignment = center if col_idx in ([1, 4] if include_images else [3]) else left

            # Embed photo
            if include_images and item.image:
                img_path = item.image.path
                if os.path.exists(img_path):
                    try:
                        xl_img = OpenpyxlImage(img_path)
                        xl_img.width  = 75
                        xl_img.height = 50
                        ws.add_image(xl_img, f'A{row_idx}')
                    except Exception:
                        pass

        # ── Freeze panes ─────────────────────────────────────────────────────
        ws.freeze_panes = 'A4'

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="Raport_Inventar_SIDESI.xlsx"'
        self._log_audit('inventory_export_excel', {
            'count': queryset.count(),
            'include_images': include_images,
        })
        return response

    # ─────────────────────────────────────────────────────────────────────────
    #   PDF EXPORT
    # ─────────────────────────────────────────────────────────────────────────
    @action(detail=False, methods=['get'], url_path='export-pdf')
    def export_pdf(self, request):
        """
        Export inventory as a premium branded PDF report.
        Query params:
          - assigned_to (UUID): filter by user
          - include_images (bool): embed item thumbnails
        """
        queryset = self.filter_queryset(self.get_queryset())
        include_images = request.query_params.get('include_images', 'false').lower() == 'true'

        buffer = BytesIO()
        page_size = landscape(A4)

        doc = SimpleDocTemplate(
            buffer,
            pagesize=page_size,
            rightMargin=1.5 * cm,
            leftMargin=1.5 * cm,
            topMargin=1.5 * cm,
            bottomMargin=1.5 * cm,
        )

        styles   = getSampleStyleSheet()
        elements = []

        # ── Custom paragraph styles ──────────────────────────────────────────
        style_company = ParagraphStyle(
            'Company',
            fontName='Helvetica-Bold',
            fontSize=18,
            textColor=SIDESI_WHITE,
            alignment=TA_LEFT,
            leading=22,
        )
        style_subtitle = ParagraphStyle(
            'Subtitle',
            fontName='Helvetica',
            fontSize=9,
            textColor=HexColor('#BFDBFE'),
            alignment=TA_LEFT,
            leading=14,
        )
        style_meta = ParagraphStyle(
            'Meta',
            fontName='Helvetica',
            fontSize=8,
            textColor=SIDESI_MUTED,
            alignment=TA_LEFT,
        )
        style_cell = ParagraphStyle(
            'Cell',
            fontName='Helvetica',
            fontSize=8,
            textColor=SIDESI_TEXT,
            leading=11,
            wordWrap='LTR',
        )
        style_cell_bold = ParagraphStyle(
            'CellBold',
            fontName='Helvetica-Bold',
            fontSize=8,
            textColor=SIDESI_TEXT,
            leading=11,
        )
        style_cell_muted = ParagraphStyle(
            'CellMuted',
            fontName='Helvetica-Oblique',
            fontSize=7,
            textColor=SIDESI_MUTED,
            leading=10,
        )
        style_header_cell = ParagraphStyle(
            'HeaderCell',
            fontName='Helvetica-Bold',
            fontSize=8,
            textColor=HexColor('#1E3A5F'),
            alignment=TA_CENTER,
            leading=11,
        )

        # ── HEADER BANNER ─────────────────────────────────────────────────────
        now_display = datetime.now().strftime('%d %B %Y, %H:%M')
        total_items = queryset.count()
        total_qty   = sum(i.quantity for i in queryset)

        header_data = [[
            Paragraph('<b>CORPORAȚIA SIDESI</b><br/>'
                      '<font size="9" color="#BFDBFE">Raport Inventar Bunuri Materiale</font>',
                      style_company),
            '',
            '',
            Paragraph(
                f'<font color="#BFDBFE" size="8">Generat: {now_display}<br/>'
                f'Total articole: {total_items} | Total cantitate: {total_qty}</font>',
                style_subtitle
            ),
        ]]

        page_w, _ = page_size
        usable_w  = page_w - 3 * cm

        header_table = Table(
            header_data,
            colWidths=[usable_w * 0.45, usable_w * 0.15, usable_w * 0.05, usable_w * 0.35],
            rowHeights=[2.2 * cm],
        )
        header_table.setStyle(TableStyle([
            ('BACKGROUND',  (0, 0), (-1, -1), SIDESI_DARK),
            ('VALIGN',      (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (0, -1),  14),
            ('RIGHTPADDING',(-1, 0),(-1, -1), 14),
            ('SPAN',        (0, 0), (2, 0)),
        ]))
        elements.append(header_table)
        elements.append(Spacer(1, 0.3 * cm))
        elements.append(HRFlowable(width='100%', thickness=2, color=SIDESI_ACCENT))
        elements.append(Spacer(1, 0.4 * cm))

        # ── DATA TABLE ────────────────────────────────────────────────────────
        if include_images:
            col_headers = ['Poză', 'Nr. Inv.', 'Denumire Obiect', 'Cant.', 'Responsabil', 'Descriere', 'Dată']
            col_widths  = [
                2.2 * cm,
                2.8 * cm,
                5.5 * cm,
                1.6 * cm,
                3.8 * cm,
                usable_w - (2.2 + 2.8 + 5.5 + 1.6 + 3.8 + 2.2) * cm,
                2.2 * cm,
            ]
        else:
            col_headers = ['Nr. Inv.', 'Denumire Obiect', 'Cant.', 'Responsabil', 'Descriere', 'Dată Alocare']
            col_widths  = [
                3.0 * cm,
                6.0 * cm,
                1.8 * cm,
                4.0 * cm,
                usable_w - (3.0 + 6.0 + 1.8 + 4.0 + 2.4) * cm,
                2.4 * cm,
            ]

        # Header row
        table_data = [[Paragraph(h, style_header_cell) for h in col_headers]]

        for row_idx, item in enumerate(queryset):
            name_p   = Paragraph(item.name, style_cell_bold)
            inv_p    = Paragraph(item.inventory_number or '—', style_cell_muted)
            desc_p   = Paragraph(item.description or '—', style_cell_muted)
            user_p   = Paragraph(_get_assigned_name(item), style_cell)
            qty_p    = Paragraph(str(item.quantity), style_cell_bold)
            date_p   = Paragraph(_format_dt(item.created_at), style_cell_muted)

            if include_images:
                if item.image and os.path.exists(item.image.path):
                    try:
                        img = RLImage(item.image.path, width=1.8 * cm, height=1.4 * cm)
                        img.hAlign = 'CENTER'
                        photo = img
                    except Exception:
                        photo = Paragraph('—', style_cell_muted)
                else:
                    photo = Paragraph('—', style_cell_muted)

                row = [photo, inv_p, name_p, qty_p, user_p, desc_p, date_p]
            else:
                row = [inv_p, name_p, qty_p, user_p, desc_p, date_p]

            table_data.append(row)

        data_table = Table(table_data, colWidths=col_widths, repeatRows=1)

        # Build alternating style commands
        style_cmds = [
            # Header
            ('BACKGROUND',   (0, 0), (-1, 0), SIDESI_MID),
            ('TEXTCOLOR',    (0, 0), (-1, 0), HexColor('#1E3A5F')),
            ('FONTNAME',     (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE',     (0, 0), (-1, 0), 8),
            ('ALIGN',        (0, 0), (-1, 0), 'CENTER'),
            ('VALIGN',       (0, 0), (-1, 0), 'MIDDLE'),
            ('BOTTOMPADDING',(0, 0), (-1, 0), 6),
            ('TOPPADDING',   (0, 0), (-1, 0), 6),
            ('LINEBELOW',    (0, 0), (-1, 0), 1.5, SIDESI_ACCENT),
            # Data rows
            ('FONTSIZE',     (0, 1), (-1, -1), 8),
            ('VALIGN',       (0, 1), (-1, -1), 'TOP'),
            ('LEFTPADDING',  (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING',   (0, 1), (-1, -1), 5),
            ('BOTTOMPADDING',(0, 1), (-1, -1), 5),
            # Grid
            ('GRID',         (0, 0), (-1, -1), 0.5, SIDESI_BORDER),
            ('LINEBELOW',    (0, -1),(-1, -1), 1, SIDESI_ACCENT),
        ]

        # Alternating row fills
        for r in range(1, len(table_data)):
            bg = SIDESI_LIGHT if r % 2 == 1 else SIDESI_WHITE
            style_cmds.append(('BACKGROUND', (0, r), (-1, r), bg))

        data_table.setStyle(TableStyle(style_cmds))
        elements.append(data_table)

        # ── FOOTER ────────────────────────────────────────────────────────────
        elements.append(Spacer(1, 0.5 * cm))
        elements.append(HRFlowable(width='100%', thickness=0.5, color=SIDESI_BORDER))
        elements.append(Spacer(1, 0.15 * cm))
        elements.append(Paragraph(
            f'<font color="#94A3B8" size="7">Corporația SIDESI  •  Raport generat automat  •  {now_display}  •  Document confidențial</font>',
            style_meta
        ))

        # ── PAGE NUMBERING ────────────────────────────────────────────────────
        def add_page_number(canvas, doc_obj):
            canvas.saveState()
            canvas.setFont('Helvetica', 7)
            canvas.setFillColor(SIDESI_MUTED)
            canvas.drawRightString(
                page_w - 1.5 * cm,
                0.8 * cm,
                f'Pagina {doc_obj.page}'
            )
            canvas.restoreState()

        doc.build(elements, onLaterPages=add_page_number, onFirstPage=add_page_number)

        buffer.seek(0)
        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="Raport_Inventar_SIDESI.pdf"'
        self._log_audit('inventory_export_pdf', {
            'count': queryset.count(),
            'include_images': include_images,
        })
        return response
